import hashlib, json
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from .config import COLLECTIONS

def load_documents(collection):
    folder=COLLECTIONS[collection]; docs=[]; failures=[]
    for path in sorted(folder.rglob('*')):
        if not path.is_file() or path.is_symlink() or path.name in {'NOTICE.txt','README.md'}: continue
        if path.suffix.lower() not in {'.pdf','.docx','.txt','.md'}: continue
        try:
            meta_file=path.with_suffix('.meta.json')
            meta=json.loads(meta_file.read_text(encoding='utf-8')) if meta_file.exists() else {}
            if path.suffix.lower()=='.pdf':
                from pypdf import PdfReader
                records=[(i,p.extract_text() or '') for i,p in enumerate(PdfReader(path).pages,1)]
            elif path.suffix.lower()=='.docx':
                from docx import Document as WordDocument
                word=WordDocument(path)
                records=[(0,'\n\n'.join([p.text for p in word.paragraphs]+[' | '.join(c.text for c in row.cells) for t in word.tables for row in t.rows]))]
            else: records=[(0,path.read_text(encoding='utf-8-sig'))]
            if not any(t.strip() for _,t in records): raise ValueError('No readable text; scan may require OCR')
            for page,text in records:
                if not text.strip(): continue
                url=meta.get('url','')
                docs.append(Document(page_content=text,metadata={'source':path.relative_to(folder).as_posix(),'title':meta.get('title',path.stem),'page':page,'url':url if url.startswith(('http://','https://')) else '', 'license':meta.get('license','User supplied; permission not established'),'collection':collection}))
        except Exception as exc: failures.append(f'{path.name}: {exc}')
    if failures: raise ValueError('Import stopped:\n'+'\n'.join(failures))
    if not docs: raise ValueError(f'No documents found in documents/{"private" if collection=="private" else "public/"+collection}')
    return docs

def split_documents(docs):
    splitter=RecursiveCharacterTextSplitter(chunk_size=1200,chunk_overlap=180,add_start_index=True)
    chunks=splitter.split_documents(docs)
    for d in chunks:
        identity=json.dumps(d.metadata,sort_keys=True)+d.page_content
        d.metadata['chunk_id']=hashlib.sha256(identity.encode()).hexdigest()
    return chunks

def snapshot_hash(chunks,model):
    content=model+'|chunk-v1|'+'|'.join(c.metadata['chunk_id'] for c in chunks)
    return hashlib.sha256(content.encode()).hexdigest()[:20]
